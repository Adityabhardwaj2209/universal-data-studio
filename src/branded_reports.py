from io import BytesIO
from datetime import datetime, timezone
from xml.sax.saxutils import escape

import pandas as pd
import streamlit as st
from PIL import Image as PILImage
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table,
    TableStyle, Image, KeepTogether,
)


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def build_branded_excel(workbook_bytes, company, title, data, logo):
    workbook = load_workbook(BytesIO(workbook_bytes))
    cover = workbook.create_sheet("Report Cover", 0)
    cover.column_dimensions["A"].width = 28
    cover.column_dimensions["B"].width = 55

    rows = [
        ["Business", company],
        ["Report", title],
        ["Period start", data["Date"].min().strftime("%Y-%m-%d")],
        ["Period end", data["Date"].max().strftime("%Y-%m-%d")],
        ["Currency", "INR"],
        ["Revenue", float(data["Revenue"].sum())],
        ["Unique orders", int(data["Order_ID"].nunique())],
        ["Units sold", float(data["Quantity"].sum())],
        ["Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")],
    ]
    for row in rows:
        cover.append([safe_cell(value) for value in row])
    for row in cover.iter_rows():
        row[0].font = Font(bold=True, color="FFFFFF")
        row[0].fill = PatternFill("solid", fgColor="8756AC")
        for cell in row:
            cell.alignment = cell.alignment.copy(wrap_text=True)
        cover.row_dimensions[row[0].row].height = 28
    cover["B6"].number_format = '"INR " #,##0.00'

    if logo:
        image = ExcelImage(BytesIO(logo))
        ratio = min(180 / image.width, 90 / image.height)
        image.width *= ratio
        image.height *= ratio
        cover.add_image(image, "D2")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_pdf(data, company, title, logo):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=(595.28, 841.89),
        rightMargin=42, leftMargin=42,
        topMargin=42, bottomMargin=48,
    )
    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.HexColor("#43205F")
    styles["Heading2"].textColor = colors.HexColor("#8756AC")
    story = []

    if logo:
        image = Image(BytesIO(logo))
        image._restrictSize(150, 65)
        story.extend([image, Spacer(1, 14)])

    story.extend([
        Paragraph(escape(company), styles["Title"]),
        Paragraph(escape(title), styles["Heading2"]),
        Paragraph(
            "Sales dates: "
            + data["Date"].min().strftime("%Y-%m-%d")
            + " to " + data["Date"].max().strftime("%Y-%m-%d"),
            styles["Normal"],
        ),
        Paragraph(
            "All monetary amounts are in INR. "
            "Includes valid rows within applied dashboard filters.",
            styles["Normal"],
        ),
        Spacer(1, 18),
    ])

    def add_table(heading, headers, rows, widths):
        story.append(Paragraph(heading, styles["Heading2"]))
        cells = [
            [Paragraph(escape(str(value)), styles["Normal"]) for value in row]
            for row in [headers] + rows
        ]
        table = Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7D6F0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F7F0FA")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#BB96CF")),
        ]))
        story.extend([table, Spacer(1, 18)])

    revenue = float(data["Revenue"].sum())
    orders = data["Order_ID"].nunique()
    add_table(
        "Sales summary",
        ["Metric", "Value"],
        [
            ["Revenue", f"{revenue:,.2f}"],
            ["Unique orders", orders],
            ["Units sold", f"{data['Quantity'].sum():,.0f}"],
            ["Average order value", f"{revenue / orders:,.2f}"],
        ],
        [255, 255],
    )

    monthly = (
        data.assign(Month=data["Date"].dt.strftime("%Y-%m"))
        .groupby("Month")["Revenue"].sum()
    )
    add_table(
        "Monthly revenue",
        ["Month", "Revenue (INR)"],
        [[month, f"{value:,.2f}"] for month, value in monthly.items()],
        [255, 255],
    )

    products = (
        data.groupby("Product")["Revenue"].sum()
        .sort_values(ascending=False).head(20)
    )
    add_table(
        "Top 20 products by revenue",
        ["Product", "Revenue (INR)"],
        [[name, f"{value:,.2f}"] for name, value in products.items()],
        [330, 180],
    )

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.setFillColor(colors.HexColor("#75617F"))
        canvas.drawString(42, 26, "Sales Studio | Generated report")
        canvas.drawRightString(553, 26, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def render_branded_exports(data, workbook_bytes):
    st.subheader("Branded reports")
    from .brand_profiles import load_brand, save_brand
    owner = st.session_state.get("account_email")
    if not owner:
        st.info("Sign in to use business branding.")
        return
    profile = load_brand(owner)
    company = st.text_input(
        "Business name",
        value=profile["company"],
        max_chars=100,
        key="brand_company_" + owner,
    ).strip()
    title = st.text_input(
        "Report title",
        value=profile["title"],
        max_chars=150,
        key="brand_title_" + owner,
    ).strip()
    uploaded_logo = st.file_uploader(
        "Business logo (optional)",
        type=["png", "jpg", "jpeg"],
        key="brand_logo_" + owner,
    )
    logo = profile["logo"]
    if uploaded_logo is not None:
        if uploaded_logo.size > 2 * 1024 * 1024:
            st.warning("Choose a logo smaller than 2 MB.")
            return
        try:
            with PILImage.open(BytesIO(uploaded_logo.getvalue())) as image:
                if image.width * image.height > 16000000:
                    raise ValueError("Logo dimensions are too large.")
                image.load()
                image.thumbnail((800, 800))
                normalized = BytesIO()
                image.convert("RGBA").save(normalized, format="PNG")
                logo = normalized.getvalue()
            st.image(logo, width=120)
        except Exception:
            st.error("Unable to read this logo. Try a smaller PNG or JPEG.")
            return


    if uploaded_logo is None and logo:
        st.caption("Saved business logo")
        st.image(logo, width=120)

    remove_logo = st.checkbox(
        "Use no logo",
        key="brand_remove_logo_" + owner,
        help="Save branding with this selected to remove the stored logo.",
    )
    if remove_logo:
        logo = None

    if st.button(
        "Save branding for future reports",
        key="brand_save_profile_" + owner,
    ):
        try:
            save_brand(owner, company, title, logo)
            st.success("Branding saved for your account.")
        except (ValueError, OSError) as error:
            st.error(str(error))

    st.caption(
        "Saved branding loads next time you sign in. "
        "Generate reports again after changing branding."
    )

    if not company or not title:
        st.info("Enter a business name and report title.")
        return

    signature = (
        company, title, logo,
        pd.util.hash_pandas_object(data, index=True).values.tobytes(),
    )
    key = "branded_ready_" + owner

    if st.button("Generate branded reports"):
        with st.spinner("Preparing PDF and Excel..."):
            try:
                st.session_state[key] = {
                    "signature": signature,
                    "pdf": build_pdf(data, company, title, logo),
                    "excel": build_branded_excel(
                        workbook_bytes, company, title, data, logo
                    ),
                }
            except Exception:
                st.error("Report generation failed. Try without the logo.")
                return

    ready = st.session_state.get(key)
    if ready and ready["signature"] == signature:
        a, b = st.columns(2)
        a.download_button(
            "Download branded PDF", ready["pdf"],
            "branded_sales_report.pdf", "application/pdf",
        )
        b.download_button(
            "Download branded Excel", ready["excel"],
            "branded_sales_report.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    elif ready:
        st.caption("Settings or data changed. Generate reports again.")
