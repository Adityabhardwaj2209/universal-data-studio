from pathlib import Path
import sys
import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
if __package__:
    from .processor import validate_sales
    from .analysis import analyze_sales
else:
    from processor import validate_sales
    from analysis import analyze_sales


def excel_safe(frame):
    frame = frame.copy()
    for column in frame.columns:
        frame[column] = frame[column].map(
            lambda value: "'" + value
            if isinstance(value, str)
            and value.lstrip().startswith(("=", "+", "-", "@"))
            else value
        )
    return frame


def create_report(source, destination="output/sales_report.xlsx"):
    clean, issues = validate_sales(source)
    summary, monthly, products, categories = analyze_sales(clean)

    summary_rows = list(summary.items())
    summary_rows.append(("Rows needing review", len(issues)))
    summary_table = pd.DataFrame(
        summary_rows, columns=["Metric", "Value"]
    )

    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)

    tables = {
        "Summary": summary_table,
        "Monthly Sales": monthly,
        "Product Sales": products,
        "Category Sales": categories,
        "Cleaned Data": clean,
        "Validation Issues": issues,
    }

    with pd.ExcelWriter(
        destination,
        engine="openpyxl",
        datetime_format="YYYY-MM-DD",
    ) as writer:
        for sheet_name, frame in tables.items():
            excel_safe(frame).to_excel(
                writer, sheet_name=sheet_name, index=False
            )

        workbook = writer.book

        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 28

            for cell in sheet[1]:
                cell.fill = PatternFill(
                    "solid", fgColor="17365D"
                )
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center")

            for column in sheet.columns:
                letter = column[0].column_letter
                width = max(
                    len(str(cell.value))
                    if cell.value is not None else 0
                    for cell in column
                )
                sheet.column_dimensions[letter].width = min(
                    max(width + 3, 14), 55
                )

            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill(
                            "solid", fgColor="EEF3F8"
                        )

            for header in sheet[1]:
                if header.value in ("Revenue", "Unit_Price"):
                    for row in sheet.iter_rows(
                        min_row=2,
                        min_col=header.column,
                        max_col=header.column,
                    ):
                        row[0].number_format = '"INR " #,##0.00'
                elif header.value == "Date":
                    for row in sheet.iter_rows(
                        min_row=2,
                        min_col=header.column,
                        max_col=header.column,
                    ):
                        row[0].number_format = "yyyy-mm-dd"

        summary_sheet = workbook["Summary"]
        for row in summary_sheet.iter_rows(min_row=2):
            if row[0].value in (
                "Total revenue", "Average order value"
            ):
                row[1].number_format = '"INR " #,##0.00'

        monthly_sheet = workbook["Monthly Sales"]
        trend = LineChart()
        trend.title = "Monthly Revenue"
        trend.y_axis.title = "Revenue (INR)"
        trend.x_axis.title = "Month"
        trend.add_data(
            Reference(
                monthly_sheet, min_col=4,
                min_row=1, max_row=monthly_sheet.max_row
            ),
            titles_from_data=True,
        )
        trend.set_categories(
            Reference(
                monthly_sheet, min_col=1,
                min_row=2, max_row=monthly_sheet.max_row
            )
        )
        trend.width = 23
        trend.height = 11
        summary_sheet.add_chart(trend, "D2")

        product_sheet = workbook["Product Sales"]
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Top 10 Products by Revenue"
        chart.x_axis.title = "Revenue (INR)"
        last_row = min(product_sheet.max_row, 11)
        chart.add_data(
            Reference(
                product_sheet, min_col=3,
                min_row=1, max_row=last_row
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                product_sheet, min_col=1,
                min_row=2, max_row=last_row
            )
        )
        chart.width = 23
        chart.height = 12
        summary_sheet.add_chart(chart, "D25")

    return destination


if __name__ == "__main__":
    source = (
        sys.argv[1] if len(sys.argv) > 1
        else "data/sample_sales.csv"
    )
    try:
        report = create_report(source)
        print("EXCEL REPORT CREATED")
        print(report.resolve())
    except PermissionError:
        print("Close sales_report.xlsx in Excel and run again.")
        sys.exit(1)
    except (ValueError, OSError) as error:
        print("ERROR:", error)
        sys.exit(1)
