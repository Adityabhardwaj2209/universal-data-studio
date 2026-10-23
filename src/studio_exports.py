from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from .studio_core import safe_excel


def excel_report(data,schema,operations,dictionary,chart_data=None):
    output=BytesIO()
    with pd.ExcelWriter(output,engine='openpyxl') as writer:
        tables={'Data':data,'Schema':schema,'History':pd.DataFrame([{'Step':i+1,'Operation':str(o)} for i,o in enumerate(operations)]),'Dictionary':pd.DataFrame(dictionary)}
        if chart_data is not None: tables['Chart values']=chart_data
        for name,frame in tables.items():
            copy=frame.copy()
            for c in copy:
                if isinstance(copy[c].dtype,pd.DatetimeTZDtype): copy[c]=copy[c].dt.tz_convert('UTC').dt.tz_localize(None)
            safe_excel(copy).to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            from openpyxl.styles import Font,PatternFill
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='8756AC')
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(42,max(14,max(len(str(c.value or '')) for c in col)+2))
    return output.getvalue()


def pdf_report(data,title,operations,chart_data=None,chart_png=None):
    output=BytesIO();styles=getSampleStyleSheet()
    font=Path(__file__).parent/'assets'/'DejaVuSans.ttf'
    if font.exists():
        if 'StudioFont' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('StudioFont',str(font)))
        for name in ['Normal','Title','Heading2']:styles[name].fontName='StudioFont'
    story=[Paragraph(escape(title),styles['Title']),Paragraph(f'{len(data):,} selected records | {len(data.columns)} fields',styles['Normal']),Spacer(1,16)]
    if chart_png:
        image=Image(BytesIO(chart_png));image._restrictSize(510,280);story.extend([image,Spacer(1,12)])
    def table(heading,frame):
        story.append(Paragraph(heading,styles['Heading2']))
        # Narrow report remains readable; complete records are in Excel.
        frame=frame.iloc[:40,:5].copy()
        cells=[[Paragraph(escape(str(c)),styles['Normal']) for c in frame.columns]]
        for row in frame.itertuples(index=False,name=None):
            cells.append([Paragraph(escape('' if pd.isna(v) else str(v)[:160]),styles['Normal']) for v in row])
        if len(frame.columns):
            t=LongTable(cells,colWidths=[510/len(frame.columns)]*len(frame.columns),repeatRows=1)
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7D6F0')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F7F0FA')]),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.extend([t,Spacer(1,14)])
    if chart_data is not None:table('Selected chart values (up to 40 groups)',chart_data)
    table('Data preview (first 40 records, first 5 fields)',data)
    story.append(Paragraph('Full data and schema are available in the Excel export. Applied transformations: '+str(len(operations)),styles['Normal']))
    def footer(canvas,doc):
        canvas.setFont('Helvetica',9);canvas.drawString(42,24,'Universal Data Studio');canvas.drawRightString(553,24,f'Page {doc.page}')
    SimpleDocTemplate(output,pagesize=(595.28,841.89),leftMargin=42,rightMargin=42,topMargin=36,bottomMargin=42).build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()
